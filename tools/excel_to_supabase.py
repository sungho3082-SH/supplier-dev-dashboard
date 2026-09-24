"""
엑셀(부품 · 일정 시트) → Supabase 가져오기용 projects.sql / projects.csv 변환기

사용법
  pip install openpyxl
  python excel_to_csv.py 개발현황.xlsx                       # 원본 그대로 CSV
  python excel_to_csv.py 개발현황.xlsx --map anonymize.json  # 가명 처리해서 CSV

만들어진 projects.sql 을 Supabase > SQL Editor 에 붙여 넣고 Run 하면 들어갑니다.
(projects.csv 는 Table Editor > Import data from CSV 용 예비 파일)

엑셀 형식 (기존 대시보드와 동일)
  [부품] Supplier | 차종 | Part No. | Part Name | Step | Open issue | Supplier Chain | All Tool | Full Tool | 공정감사 | Full CAPA
  [일정] 차종 | Proto | P1 | P2 | M | SOP | Vol.
  Open issue 칸을 여러 행에 걸쳐 병합해 두면, 그 부품들은 같은 이슈 묶음(issue_group)으로 저장됩니다.
"""
import csv, datetime, json, sys
import openpyxl

COLS = ["car_model", "supplier", "part_no", "part_name", "step", "open_issue", "supply_chain", "issue_group",
        "proto_start", "p1_start", "p2_start", "m_start", "sop_start",
        "all_tool", "full_tool", "audit", "full_capa", "volume", "sort_order"]


def as_date(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    return ""  # '미실시' 같은 글자나 빈칸은 날짜 없음


def replace_all(text, mapping):
    if not text:
        return text
    for a, b in sorted(mapping.items(), key=lambda kv: -len(kv[0])):  # 긴 이름부터 바꿔서 부분 치환 방지
        text = text.replace(a, b)
    return text


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    path = sys.argv[1]
    mapping = {}
    if "--map" in sys.argv:
        mapping = json.load(open(sys.argv[sys.argv.index("--map") + 1], encoding="utf-8"))

    wb = openpyxl.load_workbook(path)
    ws, sc = wb["부품"], wb["일정"]

    # 일정 시트: 차종 → 단계 날짜
    sched = {}
    for r in sc.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        sched[str(r[0]).strip()] = {"proto_start": as_date(r[1]), "p1_start": as_date(r[2]), "p2_start": as_date(r[3]),
                                     "m_start": as_date(r[4]), "sop_start": as_date(r[5]),
                                     "volume": str(int(r[6])) if isinstance(r[6], (int, float)) else ""}

    # Open issue(F열) 병합 범위 → 이슈 묶음
    group_of, issue_of = {}, {}
    for rng in ws.merged_cells.ranges:
        if rng.min_col == 6 and rng.max_col == 6 and rng.max_row > rng.min_row:
            gid = f"G{rng.min_row}"
            top = ws.cell(rng.min_row, 6).value
            for row in range(rng.min_row, rng.max_row + 1):
                group_of[row], issue_of[row] = gid, top

    out, order = [], 0
    for i, r in enumerate(ws.iter_rows(min_row=2, max_col=11, values_only=True), start=2):
        supplier, car = r[0], r[1]
        if not supplier or not car or not r[3]:
            continue  # 품명이 없는 메모 행은 건너뜀
        order += 1
        car = str(car).strip()
        s = sched.get(car, {})
        issue = issue_of.get(i, r[5])
        if issue in ("-", None):
            issue = ""
        row = {
            "car_model": car, "supplier": supplier, "part_no": r[2] or "", "part_name": r[3] or "",
            "step": r[4] or "", "open_issue": issue, "supply_chain": r[6] or "", "issue_group": group_of.get(i, ""),
            "proto_start": s.get("proto_start", ""), "p1_start": s.get("p1_start", ""), "p2_start": s.get("p2_start", ""),
            "m_start": s.get("m_start", ""), "sop_start": s.get("sop_start", ""),
            "all_tool": as_date(r[7]), "full_tool": as_date(r[8]), "audit": as_date(r[9]), "full_capa": as_date(r[10]),
            "volume": s.get("volume", ""), "sort_order": order,
        }
        if mapping:
            for k in ("car_model", "supplier", "part_no", "part_name", "open_issue", "supply_chain"):
                row[k] = replace_all(str(row[k]), mapping)
        out.append(row)

    with open("projects.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLS)
        w.writeheader()
        w.writerows(out)
    # 같은 내용을 SQL로도 저장 (Supabase > SQL Editor 에 붙여 넣어 실행)
    def q(v):
        if v in ("", None):
            return "null"
        if isinstance(v, int):
            return str(v)
        return "'" + str(v).replace("'", "''") + "'"
    with open("projects.sql", "w", encoding="utf-8") as f:
        f.write("insert into projects (" + ",".join(COLS) + ") values\n")
        f.write(",\n".join("(" + ",".join(q(r[c]) for c in COLS) + ")" for r in out) + ";\n")
    print(f"projects.csv / projects.sql 저장: {len(out)}행")


if __name__ == "__main__":
    main()
